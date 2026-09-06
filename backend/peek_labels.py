"""Peek at the Excel label files."""
import openpyxl
from pathlib import Path

for xlsx in [r"D:\dataset anemia\India\India.xlsx", r"D:\dataset anemia\Italy\Italy.xlsx"]:
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb.active
    print(f"\n=== {Path(xlsx).name} (sheet: {ws.title}) ===")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 5:
            print(f"  Row {i}: {row}")
        else:
            break
    print(f"  ... (total rows: {ws.max_row}, cols: {ws.max_column})")
    wb.close()
